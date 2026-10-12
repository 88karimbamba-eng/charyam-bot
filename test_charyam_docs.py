"""Tests de charyam_docs : python -I -m unittest test_charyam_docs -v

Aucun réseau, aucune clé. Les tests Word/PDF vérifient la structure et le contenu
relu par python-docx / pypdf, pas le rendu dans Microsoft Word.
"""
import io
import os
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile
from decimal import Decimal

import charyam_docs as cd

try:
    import docx  # python-docx, uniquement pour relire
except ImportError:  # pragma: no cover
    docx = None
try:
    import pypdf
    import reportlab  # noqa: F401
except ImportError:  # pragma: no cover
    pypdf = None

BLOCKS = [
    {"type": "h1", "text": "Expérience professionnelle"},
    {"type": "p", "text": "Chargé de portefeuille — œuvre à l'état d'équilibre, café & thé <ok>"},
    {"type": "bullet", "text": "Suivi de 500 producteurs"},
    {"type": "num", "n": 2, "text": "Recouvrement"},
    {"type": "table", "rows": [["Poste", "Période"], ["Chargé", "2019-2026"]]},
]


class TestMarkdown(unittest.TestCase):
    def test_structure(self):
        md = ("# Titre\n\nUne ligne\nsuite.\n\n## Sous-titre\n- a **gras**\n* b\n1. un\n2) deux\n\n"
              "| A | B |\n|---|:--:|\n| 1 | 2 |\n| 3 |\n\nFin")
        b = cd.markdown_to_blocks(md)
        self.assertEqual([x["type"] for x in b], ["h1", "p", "h2", "bullet", "bullet", "num", "num", "table", "p"])
        self.assertEqual(b[1]["text"], "Une ligne suite.")
        self.assertEqual(b[3]["text"], "a gras")
        self.assertEqual(b[7]["rows"], [["A", "B"], ["1", "2"], ["3", ""]])  # séparateur retiré, ligne complétée

    def test_controle_et_vide(self):
        self.assertEqual(cd.markdown_to_blocks(""), [])
        b = cd.markdown_to_blocks("a\x00b\x0bc")
        self.assertEqual(b[0]["text"], "abc")


class TestDocx(unittest.TestCase):
    def test_xml_bien_forme(self):
        data = cd.docx_bytes("Mon CV & <test>", BLOCKS)
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            self.assertIsNone(z.testzip())
            for name in z.namelist():
                ET.fromstring(z.read(name))  # lève si mal formé
            self.assertIn("word/document.xml", z.namelist())

    @unittest.skipIf(docx is None, "python-docx absent")
    def test_relecture_python_docx(self):
        d = docx.Document(io.BytesIO(cd.docx_bytes("Titre éàç", BLOCKS)))
        texts = [p.text for p in d.paragraphs]
        self.assertIn("Titre éàç", texts)
        self.assertIn("Expérience professionnelle", texts)
        self.assertTrue(any("œuvre" in t and "<ok>" in t for t in texts))  # caractères spéciaux conservés
        self.assertEqual(d.tables[0].rows[1].cells[1].text, "2019-2026")
        self.assertEqual(d.paragraphs[0].style.name, "Title")

    def test_injection_et_controles(self):
        data = cd.docx_bytes("t\x00", [{"type": "p", "text": "</w:t><w:evil/>\x01\ud800 ok"}])
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            root = ET.fromstring(z.read("word/document.xml"))
        self.assertNotIn("evil", [e.tag.split("}")[1] for e in root.iter()])

    def test_gros_volume(self):
        blocks = [{"type": "p", "text": "x" * 500}] * 3000  # plafonné à MAX_BLOCKS
        self.assertLess(len(cd.docx_bytes("t", blocks)), 5_000_000)


@unittest.skipIf(pypdf is None, "reportlab/pypdf absents")
class TestPdf(unittest.TestCase):
    def test_contenu(self):
        data = cd.pdf_bytes("Rapport de visite", BLOCKS)
        self.assertTrue(data.startswith(b"%PDF"))
        r = pypdf.PdfReader(io.BytesIO(data))
        text = "".join(p.extract_text() for p in r.pages)
        for needle in ("Rapport de visite", "Expérience professionnelle", "œuvre", "<ok>", "2019-2026"):
            self.assertIn(needle, text)

    def test_emoji_ne_plante_pas(self):
        data = cd.pdf_bytes("t", [{"type": "p", "text": "Bravo 🎉 & merci"}])
        self.assertTrue(data.startswith(b"%PDF"))

    def test_plusieurs_pages(self):
        blocks = [{"type": "p", "text": "ligne " + str(i)} for i in range(400)]
        self.assertGreater(len(pypdf.PdfReader(io.BytesIO(cd.pdf_bytes("t", blocks))).pages), 1)


class TestDevis(unittest.TestCase):
    LINES = [
        {"designation": "Visite terrain", "quantite": "3", "prix_unitaire": "125 000"},
        {"designation": "Rapport", "quantite": "2,5", "prix_unitaire": 40000},
    ]

    def test_totaux(self):
        q = cd.compute_quote(self.LINES, tva_rate="18")
        self.assertEqual([l["total"] for l in q["lignes"]], [375000, 100000])
        self.assertEqual((q["total_ht"], q["tva"], q["total_ttc"]), (475000, 85500, 560500))

    def test_sans_tva(self):
        q = cd.compute_quote(self.LINES)
        self.assertEqual((q["tva"], q["total_ttc"]), (0, 475000))

    def test_arrondi_moitie_haut(self):
        q = cd.compute_quote([{"designation": "x", "quantite": 3, "prix_unitaire": "333,5"}])
        self.assertEqual(q["total_ht"], 1001)  # 1000,5 -> 1001

    def test_pas_de_derive_flottante(self):
        # Decimal : 1000 x 1.0005 = 1000.5 exactement -> 1001 (moitié vers le haut)
        q = cd.compute_quote([{"designation": "x", "quantite": "1000", "prix_unitaire": "1.0005"}])
        self.assertEqual(q["total_ht"], 1001)

    def test_arrondi_par_ligne_documente(self):
        # Chaque ligne vaut 0,3 FCFA -> arrondie à 0 ; le total est la somme des lignes arrondies.
        q = cd.compute_quote([{"designation": "x", "quantite": "3", "prix_unitaire": "0.1"}] * 10)
        self.assertEqual(q["total_ht"], 0)

    def test_refus(self):
        bad = [
            ([], "0"),
            ([{"designation": "x", "quantite": "-1", "prix_unitaire": "10"}], "0"),
            ([{"designation": "x", "quantite": "0", "prix_unitaire": "10"}], "0"),
            ([{"designation": "x", "quantite": "1", "prix_unitaire": "-5"}], "0"),
            ([{"designation": "x", "quantite": "abc", "prix_unitaire": "10"}], "0"),
            ([{"designation": "x", "quantite": "NaN", "prix_unitaire": "10"}], "0"),
            ([{"designation": "x", "quantite": "Infinity", "prix_unitaire": "10"}], "0"),
            ([{"designation": "x", "quantite": "1", "prix_unitaire": "1.250.000"}], "0"),
            ([{"designation": "x", "quantite": "1", "prix_unitaire": "10"}], "150"),
            ([{"designation": "x", "prix_unitaire": "10"}], "0"),
        ]
        for lines, tva in bad:
            with self.subTest(lines=lines, tva=tva), self.assertRaises(ValueError):
                cd.compute_quote(lines, tva)

    @unittest.skipIf(pypdf is None, "reportlab/pypdf absents")
    def test_devis_vers_pdf_et_docx(self):
        q = cd.compute_quote(self.LINES, "18")
        blocks = cd.quote_blocks(q, objet="Mission", client="Coopérative X",
                                 donnees_client=["3 visites demandées"], hypotheses=["TVA 18 %"])
        text = "".join(p.extract_text() for p in pypdf.PdfReader(io.BytesIO(cd.pdf_bytes("Devis", blocks))).pages)
        self.assertIn("560", text)  # total TTC 560 500, séparateur de milliers variable
        self.assertIn("Hypothèses", text)
        if docx:
            d = docx.Document(io.BytesIO(cd.docx_bytes("Devis", blocks)))
            last = d.tables[0].rows[-1].cells
            self.assertEqual(last[1].text, "Total TTC")
            self.assertEqual(last[4].text.replace(" ", ""), "560500")


class TestFichiersEtIsolation(unittest.TestCase):
    def test_noms(self):
        self.assertEqual(cd.safe_filename("CV Kouamé/../etc?", "docx"), "CV_Kouame_etc.docx")
        self.assertEqual(cd.safe_filename("", ".pdf"), "document.pdf")
        self.assertLessEqual(len(cd.safe_filename("a" * 500, "pdf")), 64)

    def test_aucun_fichier_cree_et_sorties_independantes(self):
        with tempfile.TemporaryDirectory() as tmp:
            old = os.getcwd()
            os.chdir(tmp)
            try:
                a = cd.docx_bytes("Alice", [{"type": "p", "text": "secret-alice"}])
                b = cd.docx_bytes("Bob", [{"type": "p", "text": "secret-bob"}])
                if pypdf:
                    cd.pdf_bytes("Alice", [{"type": "p", "text": "x"}])
                self.assertEqual(os.listdir(tmp), [])
            finally:
                os.chdir(old)
        self.assertNotIn(b"secret-alice", zipfile.ZipFile(io.BytesIO(b)).read("word/document.xml"))
        self.assertNotIn(b"secret-bob", zipfile.ZipFile(io.BytesIO(a)).read("word/document.xml"))


if __name__ == "__main__":
    unittest.main()
